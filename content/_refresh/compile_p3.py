# -*- coding: utf-8 -*-
import json, glob, os, re
from openpyxl import load_workbook
ANS=r'C:/Users/joysh/caspr/caspr-claude-core/content/_refresh/answers'
XL=r'C:/Users/joysh/caspr/caspr-claude-core/content/caspr-help-content-repository.xlsx'

updated={}
for fp in sorted(glob.glob(os.path.join(ANS,'p3_*.json'))):
    for it in json.load(open(fp,encoding='utf-8')):
        a=it.get('answer') or it.get('answer_l2')
        if a: updated[it['id']]=a.strip()
    print(os.path.basename(fp),'->',len(json.load(open(fp,encoding='utf-8'))))
print('Collected updated answers:',len(updated))

# QA on updated set
BANNED=['platform','leverage','algorithm','workflow','revolutionary','game-chang','powerful ai','excited to announce']
excl=[];banned=[];one_m=[];reset_acct=[];lc_dataroom=[];menu_k=[]
for cid,a in updated.items():
    low=a.lower()
    if '!' in a: excl.append(cid)
    for b in BANNED:
        if b in low: banned.append((cid,b))
    if re.search(r'\b1m\+|1 million\b', low): one_m.append(cid)
    if 'reset account' in low or 'reset your account' in low: reset_acct.append(cid)
    if 'data room' in a and re.search(r'(?<![A-Za-z])data room', a):  # lowercase 'data room'
        if re.search(r'(?<![>\w])data room', a) and 'Data Room' not in a.split('data room')[0][-40:]:
            pass
    lc = len(re.findall(r'(?<![A-Za-z>])data room', a))
    if lc: lc_dataroom.append((cid,lc))
    if cid.startswith('K-') and re.search(r"(its menu|the menu|file's menu|row menu|\u2261 menu)", low):
        menu_k.append(cid)
print('Exclamation:',excl or 'none')
print('Banned:',banned or 'none')
print('Remaining 1M+/1 million:',one_m or 'none (all -> 25M+)')
print('Still says reset (your) account:',reset_acct or 'none')
print('Lowercase "data room" occurrences:',lc_dataroom or 'none')
print('K cards still referencing a menu:',menu_k or 'none')

# write to workbook
wb=load_workbook(XL); ac=wb['Answer Cards']
h={ac.cell(row=1,column=c).value:c for c in range(1,ac.max_column+1)}
CID=h['ID'];CAN=h['Canonical Answer (writer fills)'];STA=h['Status'];OWN=h['Owner']
idx={ac.cell(row=r,column=CID).value:r for r in range(3,ac.max_row+1)}
n=0
for cid,a in updated.items():
    r=idx.get(cid)
    if r: ac.cell(row=r,column=CAN,value=a); ac.cell(row=r,column=STA,value='Draft'); ac.cell(row=r,column=OWN,value='AI draft (pass3)'); n+=1
wb.save(XL)
print('Written to workbook:',n)
expected=108  # all except E-01..E-05
print('Expected 108 (E excluded). Got',len(updated),'->', 'OK' if len(updated)==expected else 'CHECK')
