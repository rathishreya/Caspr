# -*- coding: utf-8 -*-
import json
from openpyxl import load_workbook
XL=r'C:/Users/joysh/caspr/caspr-claude-core/content/caspr-help-content-repository.xlsx'
OUT=r'C:/Users/joysh/caspr/caspr-claude-core/content/_refresh/current_answers.json'
wb=load_workbook(XL); ac=wb['Answer Cards']
h={ac.cell(row=1,column=c).value:c for c in range(1,ac.max_column+1)}
out={}
for r in range(3,ac.max_row+1):
    cid=ac.cell(row=r,column=h['ID']).value
    if not cid or cid=='EX-00': continue
    a=ac.cell(row=r,column=h['Canonical Answer (writer fills)']).value
    out[cid]={'id':cid,'question':ac.cell(row=r,column=h['Question (customer words)']).value,
              'depth':ac.cell(row=r,column=h['Max Depth']).value,
              'word_count':ac.cell(row=r,column=h['Word Count']).value,
              'answer':a or ''}
json.dump(out,open(OUT,'w',encoding='utf-8'),ensure_ascii=False,indent=1)
print('emitted',len(out),'cards')
